"""docs/sample_loans.xlsx for non-coders: 10 loans of the TRAIN split (5 honest incl. a livestock seller, a cash-heavy business and a
generic-trade supplier; 5 fraud: T1, T2, T3, T4 and one stealth case) with their transactions and a short Uzbek explanation per typology.

    .venv/Scripts/python -m ml.eval.make_sample_xlsx

Synthetic data only (generator seed 42); names are masked. The loans are picked deterministically (fixed seed) among loans with a readable
history of 6-14 transactions.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ml.generator.build import build
from ml.generator.config import GENERIC_OKED, SPLITS

OUT = Path(__file__).resolve().parents[2] / "docs" / "sample_loans.xlsx"
SEED = 42
TYPE_UZ = {"disbursement": "Kredit tushdi", "cash_withdrawal": "Naqd pul yechildi", "supplier_payment": "Ta'minotchiga to'lov",
           "p2p_out": "Shaxsga o'tkazma", "p2p_in": "Shaxsdan kirim", "external_transfer": "Boshqa bankka o'tkazma", "income": "Tushum"}
KIND_UZ = {"person": "shaxs", "business": "tashkilot", "external": "tashqi hisob", "supplier": "ta'minotchi", "none": "-"}
PURPOSE_UZ = {"tikuv_dastgohi": "tikuv dastgohi", "chorva": "chorva", "issiqxona": "issiqxona", "parranda": "parranda",
              "asalarichilik": "asalarichilik"}

EXPLAIN = [
    ("halol", "Halol kredit",
     "Qarz oluvchi kredit mablag'ini ta'minotchiga o'tkazadi, qolgan qismini kichik xarajatlar uchun ishlatadi va keyingi oylarda muntazam tushum oladi. "
     "Halol kreditlar orasida ham 'shubhali ko'rinadigan' holatlar bor: qonuniy ravishda chorva yoki parranda sotgan fermer (bir martalik katta kirim va naqd yechish), "
     "naqd bilan ishlaydigan qishloq biznesi (ish haqi naqd to'lanadi) va umumiy savdo (OKED mos kelmaydigan) ta'minotchidan qonuniy xarid. "
     "Tizim ularni 'aybdor' demaydi, faqat 'avval shu kreditni tekshiring, sababi shu' deydi; yakuniy qarorni bank xodimi qabul qiladi."),
    ("T1", "T1: tez naqdlashtirish",
     "Kredit tushgach, mablag'ning katta qismi birinchi 1-2 kunda naqd yechiladi va ta'minotchiga to'lov deyarli qilinmaydi. "
     "Yashirin (stealth) variantda pul 5-10 kun davomida bo'lib-bo'lib yechiladi va ta'minotchiga kichik to'lov o'tkaziladi. "
     "Asos: brief §2, xodim yaqinining xo'jaligi orqali 71 fuqaroning kreditlari naqdlashtirilgan va soxta ish haqi ma'lumotlari bilan mablag' o'zlashtirilgan holatlar."),
    ("T2", "T2: ta'minotchi halqasi",
     "Ko'p qarz oluvchi pulni bitta ta'minotchiga o'tkazadi, ta'minotchi esa uni tez naqdlashtiradi yoki boshqa hisobga chiqaradi; ta'minotchi bank xodimi bilan bog'langan "
     "bo'lishi mumkin. Alohida qarz oluvchining o'z hisobi oddiy ko'rinadi: signal ta'minotchi darajasida (nechta kredit, qancha tez naqd chiqarilgan, qancha davr faol). "
     "Asos: brief §2, bank xodimi o'z yaqinining xo'jaligini ta'minotchi qilib, 71 fuqaroning 2.18 mlrd so'mlik kreditini naqdlashtirgan."),
    ("T3", "T3: aktivni sotib yuborish",
     "Ta'minotchiga to'lov bor (aktiv sotib olingan), keyin bir martalik katta kirim keladi (aktiv sotilgan) va u naqd yechiladi yoki boshqa hisobga o'tadi; muntazam tushum ko'pincha to'xtaydi. "
     "Kredit yaqinda berilgan bo'lsa, bu hodisa hali sodir bo'lmagan bo'lishi mumkin. Asos: brief §2, 160 bosh naslli chorva naqdga sotib yuborilgan; tikuv dastgohlari naqdga sotilgan."),
    ("T4", "T4: soxta hisob-faktura",
     "Ta'minotchining faoliyat turi (OKED) kredit maqsadiga mos kelmaydi yoki hisob-faktura bo'yicha to'lov qilinmagan; pul oqimi o'zi oddiy ko'rinishi mumkin. "
     "Halol kreditlar orasida ham OKED mos kelmaydigan ta'minotchilar bor (umumiy savdo), shuning uchun bu belgi yolg'iz o'zi isbot emas. "
     "Asos: brief §2, uskuna olib kelinmasdan, soxta hujjatlar bilan 'ishlab chiqarish boshlandi' deyilgan (8 mlrd so'm)."),
    ("eslatma", "Muhim eslatma",
     "Bu ma'lumotlar SINTETIK (generator, seed 42): haqiqiy odamlar va haqiqiy bank ma'lumotlari yo'q, ismlar maskalangan. Tur ustuni faqat generator yorlig'i; haqiqiy tizim "
     "bu yorliqni ko'rmaydi va hech kimni 'aybdor' demaydi. Aniqlik ko'rsatkichlari sintetik ma'lumotlarda o'lchangan; haqiqiy aniqlik pilotda tarixiy kreditlarda o'lchanadi."),
]


def pick(rng, ids):
    ids = sorted(ids)
    return ids[int(rng.integers(len(ids)))] if ids else None


def main():
    t = build(SPLITS["train"])
    loans, tx, lab, sup = t["loans"], t["transactions"], t["labels"].set_index("loan_id"), t["suppliers"].set_index("supplier_id")
    n_tx = tx[tx.loan_id.notna()].groupby("loan_id").size()
    readable = set(n_tx[(n_tx >= 6) & (n_tx <= 14)].index)
    L = loans.set_index("loan_id")
    generic = {i for i in lab.index if L.at[i, "supplier_id"] in set(sup[sup.oked_code.isin(GENERIC_OKED)].index)}
    rng = np.random.default_rng(SEED)
    chosen: list[tuple[str, str, str]] = []                 # (loan_id, turi, izoh)
    used: set[str] = set()

    def add(ids, turi, note):
        i = pick(rng, set(ids) - used)
        used.add(i)
        chosen.append((i, turi, note))

    honest = ~lab.is_fraud & lab.index.isin(readable)
    add(lab.index[honest & (lab.profile == "livestock_seller")], "halol", "qonuniy ravishda chorva yoki parranda sotgan fermer")
    add(lab.index[honest & (lab.profile == "cash_heavy")], "halol", "naqd bilan ishlaydigan qishloq biznesi")
    add(lab.index[honest & (lab.profile == "oked_mismatch") & lab.index.isin(generic)], "halol", "umumiy savdo (OKED mos kelmaydi) ta'minotchisidan qonuniy xarid")
    add(lab.index[honest & (lab.profile == "normal")], "halol", "oddiy halol kredit")
    add(lab.index[honest & (lab.profile == "normal")], "halol", "oddiy halol kredit")
    vis = lab.is_fraud & lab.fraud_visible & lab.index.isin(readable)
    add(lab.index[vis & (lab.typology == "T1") & ~lab.stealth], "T1", "tez naqdlashtirish")
    add(lab.index[vis & (lab.typology == "T2") & ~lab.stealth], "T2", "ta'minotchi halqasi")
    add(lab.index[vis & (lab.typology == "T3")], "T3", "aktivni sotib yuborish")
    add(lab.index[vis & (lab.typology == "T4") & ~lab.stealth], "T4", "soxta hisob-faktura (OKED mos kelmaydi)")
    add(lab.index[vis & (lab.typology == "T1") & lab.stealth], "T1", "yashirin (stealth): naqd pul 5-10 kun davomida bo'lib-bo'lib yechilgan")
    ids = [c[0] for c in chosen]
    order = {i: k for k, i in enumerate(ids)}

    wb = Workbook()
    ws = wb.active
    ws.title = "Kreditlar"
    head = ["Kredit ID", "Qarz oluvchi (maska)", "Maqsad", "Summa (so'm)", "Muddat (oy)", "Berilgan sana", "Hudud", "Filial",
            "Ta'minotchi ID", "Ta'minotchi nomi", "Ta'minotchi OKED", "Hisob-faktura summasi (so'm)", "Turi", "Izoh"]
    ws.append(head)
    for lid, turi, note in chosen:
        r, s = L.loc[lid], sup.loc[L.at[lid, "supplier_id"]]
        ws.append([lid, r.borrower_name_masked, PURPOSE_UZ[str(r.purpose)], int(r.amount_uzs), int(r.term_months),
                   r.disbursed_at.strftime("%Y-%m-%d"), str(r.region), r.branch_id, r.supplier_id, s["name"], s["oked_code"],
                   int(r.invoice_amount_uzs), turi, note])
    ws2 = wb.create_sheet("Tranzaksiyalar")
    ws2.append(["Kredit ID", "Turi", "Sana va vaqt", "Kredit tushganidan keyin (kun)", "Operatsiya", "Summa (so'm)", "Kontragent turi"])
    d = tx[tx.loan_id.isin(ids)].copy()
    d["k"] = d.loan_id.map(order)
    d = d.sort_values(["k", "ts", "txn_id"])
    turi = {i: t_ for i, t_, _ in chosen}
    for _, r in d.iterrows():
        day = (r.ts - L.at[r.loan_id, "disbursed_at"]) / pd.Timedelta(days=1)
        ws2.append([r.loan_id, turi[r.loan_id], r.ts.strftime("%Y-%m-%d %H:%M"), round(float(day), 1), TYPE_UZ[str(r["type"])],
                    int(r.amount_uzs), KIND_UZ[str(r.counterparty_kind)]])
    ws3 = wb.create_sheet("Turlar tavsifi")
    ws3.append(["Tur", "Nomi", "Tavsif"])
    for row in EXPLAIN:
        ws3.append(list(row))

    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    for sheet, widths in ((ws, [11, 22, 16, 15, 11, 14, 20, 9, 14, 28, 12, 22, 8, 55]), (ws2, [11, 8, 18, 18, 24, 15, 16]), (ws3, [10, 28, 140])):
        for c in sheet[1]:
            c.font, c.fill = bold, fill
        for k, w in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(k)].width = w
        sheet.freeze_panes = "A2"
    for row in ws3.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for sheet in (ws, ws2):
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, int) and c.value > 9999:
                    c.number_format = "#,##0"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT, len(ids), "loans,", len(d), "transactions")


if __name__ == "__main__":
    main()
